
import re
import sys
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ============================================================
# CONFIGURATION
# ============================================================

BASE_DIR = Path.cwd()
OUTPUT_DIR = BASE_DIR / "output"
MASTER_FILE = OUTPUT_DIR / "Appeal_Tracker_Master.xlsx"

RUN_TIME = datetime.now()
RUN_TIMESTAMP = RUN_TIME.strftime("%Y-%m-%d %H:%M:%S")
RUN_ID = (
    sys.argv[2]
    if len(sys.argv) > 2
    else RUN_TIME.strftime("%Y%m%d_%H%M%S")
)

LOG_FILE = (
    Path(sys.argv[1])
    if len(sys.argv) > 1
    else OUTPUT_DIR / "appeal_run.log"
)

if not LOG_FILE.is_absolute():
    LOG_FILE = BASE_DIR / LOG_FILE

RUN_REPORT = OUTPUT_DIR / f"Appeal_Tracker_{RUN_ID}.xlsx"

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

if not LOG_FILE.exists():
    raise SystemExit(f"ERROR: Log file not found: {LOG_FILE}")

log_text = LOG_FILE.read_text(errors="replace")

# ============================================================
# PARSE CLAIM-LEVEL LOG SECTIONS
# ============================================================

starts = list(re.finditer(
    r"Generating\s+(Aetna|UHC|BCBS)\s+PDF",
    log_text,
    re.IGNORECASE
))

if not starts:
    raise SystemExit(
        "ERROR: No claim sections found. Check your run log format."
    )

claims = []

for i, match in enumerate(starts):
    block_start = match.start()
    block_end = (
        starts[i + 1].start()
        if i + 1 < len(starts)
        else len(log_text)
    )

    block = log_text[block_start:block_end].strip()
    payer = match.group(1).upper()

    def get_value(pattern, default=""):
        result = re.search(
            pattern,
            block,
            re.IGNORECASE | re.MULTILINE
        )
        return result.group(1).strip() if result else default

    # Claim ID
    claim_id = get_value(
        r"Claim(?:\s+ID)?\s*[:#]?\s*(\d+)"
    )

    if not claim_id:
        result = re.search(r"\b(500\d{3})\b", block)
        claim_id = (
            result.group(1)
            if result
            else f"UNKNOWN-{i + 1}"
        )

    # Claim information
    denial = get_value(
        r"Denial(?:\s+Category)?\s*:\s*(.+)",
        "NOT FOUND"
    )

    folder = get_value(
        r"Document Folder\s*:\s*(.+)",
        ""
    )

    pdf_path_text = get_value(
        r"(?:Created|Saved|Generated)(?:\s+PDF)?\s*:\s*(.+\.pdf)",
        ""
    )

    # Attachment count
    attachment_count = None

    attachment_patterns = [
        r"(\d+)\s+attachments?\s+(?:added|attached)",
        r"attached\s+(\d+)\s+files?",
        r"attachments?\s*:\s*(\d+)",
        r"(\d+)\s+files?\s+attached",
    ]

    for pattern in attachment_patterns:
        result = re.search(pattern, block, re.IGNORECASE)
        if result:
            attachment_count = int(result.group(1))
            break

    # Resolve paths
    folder_path = Path(folder) if folder else None
    pdf_path = Path(pdf_path_text) if pdf_path_text else None

    if folder_path and not folder_path.is_absolute():
        folder_path = BASE_DIR / folder_path

    if pdf_path and not pdf_path.is_absolute():
        pdf_path = BASE_DIR / pdf_path

    # Check document folder
    folder_exists = bool(
        folder_path and folder_path.is_dir()
    )

    # Check generated PDF
    pdf_exists = bool(
        pdf_path and pdf_path.is_file()
    )

    # ========================================================
    # STATUS RULE REQUESTED BY USER
    # Only missing document folders trigger manual review.
    # ========================================================

    if not folder_exists:
        status = "MANUAL REVIEW"
        validation_issues = "Document folder missing or not found"
    else:
        status = "PASS - READY FOR FINAL QA"
        validation_issues = "None"

    # Keep parsing warnings visible without changing status.
    if re.search(
        r"Ignoring wrong pointing object",
        block,
        re.IGNORECASE
    ):
        pdf_warnings = "PDF parsing warning - inspect PDF"
    else:
        pdf_warnings = "None detected"

    claims.append({
        "Run ID": RUN_ID,
        "Run Timestamp": RUN_TIMESTAMP,
        "Claim ID": claim_id,
        "Payer": payer,
        "Denial Category": denial,
        "Document Folder": str(folder_path) if folder_path else folder,
        "Folder Exists": "Yes" if folder_exists else "No",
        "Attachment Count": (
            attachment_count
            if attachment_count is not None
            else "Not found in log"
        ),
        "PDF Path": str(pdf_path) if pdf_path else pdf_path_text,
        "PDF Exists": "Yes" if pdf_exists else "No",
        "Validation Issues": validation_issues,
        "PDF Warnings": pdf_warnings,
        "Final Status": status,
        "Run Log Details": block,
        "Reviewer Notes": "",
    })

# ============================================================
# COUNTS FOR THIS RUN
# ============================================================

total_claims = len(claims)

manual_claims = [
    c for c in claims
    if c["Final Status"] == "MANUAL REVIEW"
]

passed_claims = [
    c for c in claims
    if c["Final Status"] == "PASS - READY FOR FINAL QA"
]

missing_folders = sum(
    c["Folder Exists"] == "No"
    for c in claims
)

missing_pdfs = sum(
    c["PDF Exists"] == "No"
    for c in claims
)

# ============================================================
# EXCEL HEADERS
# ============================================================

claim_headers = list(claims[0].keys())

review_headers = [
    "Run ID",
    "Run Timestamp",
    "Claim ID",
    "Payer",
    "Denial Category",
    "Document Folder",
    "Validation Issues",
    "Run Log Details",
    "Reviewer Notes",
]

history_headers = [
    "Run ID",
    "Run Timestamp",
    "Log File",
    "Total Claims",
    "Passed Claims",
    "Manual Review",
    "Missing PDFs",
    "Missing Folders",
    "Individual Report",
]

# ============================================================
# EXCEL FORMATTING
# ============================================================

def format_workbook(wb):
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"

        if ws.max_row >= 2:
            ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )
            cell.fill = PatternFill(
                "solid",
                fgColor="1F4E78"
            )
            cell.alignment = Alignment(
                wrap_text=True,
                vertical="top"
            )

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    wrap_text=True,
                    vertical="top"
                )

        for column in ws.columns:
            letter = get_column_letter(column[0].column)

            longest = max(
                (
                    len(str(cell.value))
                    for cell in column
                    if cell.value is not None
                ),
                default=12
            )

            ws.column_dimensions[letter].width = min(
                max(longest + 2, 14),
                40
            )

        # Make the full run-log column readable.
        for cell in ws[1]:
            if cell.value == "Run Log Details":
                ws.column_dimensions[
                    get_column_letter(cell.column)
                ].width = 70

    if "Summary" in wb.sheetnames:
        wb["Summary"].column_dimensions["A"].width = 38
        wb["Summary"].column_dimensions["B"].width = 90

# ============================================================
# CREATE INDIVIDUAL REPORT FOR THIS RUN
# ============================================================

wb_run = Workbook()

summary = wb_run.active
summary.title = "Summary"

runlog = wb_run.create_sheet("Claim-Level Run Log")
review = wb_run.create_sheet("Manual Review")

summary.append(["Appeal Automation Run Summary"])
summary.append(["Run ID", RUN_ID])
summary.append(["Run Timestamp", RUN_TIMESTAMP])
summary.append(["Total Claims", total_claims])
summary.append(["Passed Claims", len(passed_claims)])
summary.append(["Manual Review", len(manual_claims)])
summary.append(["Missing Folders", missing_folders])
summary.append(["Missing PDFs", missing_pdfs])
summary.append([])
summary.append([
    "Status Rule",
    "Folder exists = PASS; folder missing = MANUAL REVIEW."
])
summary.append([
    "Important",
    "PASS only reflects folder existence. Verify required files, "
    "attachment content, claim information and PDF before submission."
])

runlog.append(claim_headers)

for claim in claims:
    runlog.append([
        claim[h] for h in claim_headers
    ])

review.append(review_headers)

for claim in manual_claims:
    review.append([
        claim["Run ID"],
        claim["Run Timestamp"],
        claim["Claim ID"],
        claim["Payer"],
        claim["Denial Category"],
        claim["Document Folder"],
        claim["Validation Issues"],
        claim["Run Log Details"],
        "",
    ])

format_workbook(wb_run)
wb_run.save(RUN_REPORT)

# ============================================================
# UPDATE MASTER WORKBOOK WITHOUT DELETING OLD RUNS
# ============================================================

if MASTER_FILE.exists():
    wb_master = load_workbook(MASTER_FILE)
else:
    wb_master = Workbook()
    ws = wb_master.active
    ws.title = "Run History"
    ws.append(history_headers)

    ws = wb_master.create_sheet("Claim-Level Run Log")
    ws.append(claim_headers)

    ws = wb_master.create_sheet("Manual Review")
    ws.append(review_headers)

# Ensure all required sheets exist
if "Run History" not in wb_master.sheetnames:
    ws = wb_master.create_sheet("Run History", 0)
    ws.append(history_headers)

if "Claim-Level Run Log" not in wb_master.sheetnames:
    ws = wb_master.create_sheet("Claim-Level Run Log")
    ws.append(claim_headers)

if "Manual Review" not in wb_master.sheetnames:
    ws = wb_master.create_sheet("Manual Review")
    ws.append(review_headers)

# Avoid appending the same run twice
history_ws = wb_master["Run History"]

existing_run_ids = {
    str(row[0])
    for row in history_ws.iter_rows(
        min_row=2,
        values_only=True
    )
    if row and row[0] is not None
}

if RUN_ID in existing_run_ids:
    raise SystemExit(
        f"Run ID {RUN_ID} already exists in the master workbook. "
        "Use a new run ID."
    )

# Append run summary
history_ws.append([
    RUN_ID,
    RUN_TIMESTAMP,
    str(LOG_FILE),
    total_claims,
    len(passed_claims),
    len(manual_claims),
    missing_pdfs,
    missing_folders,
    str(RUN_REPORT),
])

# Append every claim from this run
master_log = wb_master["Claim-Level Run Log"]

for claim in claims:
    master_log.append([
        claim[h] for h in claim_headers
    ])

# Append only folder-missing claims to the review history
master_review = wb_master["Manual Review"]

for claim in manual_claims:
    master_review.append([
        claim["Run ID"],
        claim["Run Timestamp"],
        claim["Claim ID"],
        claim["Payer"],
        claim["Denial Category"],
        claim["Document Folder"],
        claim["Validation Issues"],
        claim["Run Log Details"],
        "",
    ])

format_workbook(wb_master)
wb_master.save(MASTER_FILE)

# ============================================================
# PRINT RESULTS
# ============================================================

print("\nAppeal tracker generated successfully.")
print(f"Run ID: {RUN_ID}")
print(f"Run timestamp: {RUN_TIMESTAMP}")
print(f"Total claims: {total_claims}")
print(f"Passed claims: {len(passed_claims)}")
print(f"Manual review: {len(manual_claims)}")
print(f"Missing folders: {missing_folders}")
print(f"Missing PDFs: {missing_pdfs}")
print(f"Individual report: {RUN_REPORT}")
print(f"Master workbook: {MASTER_FILE}")
